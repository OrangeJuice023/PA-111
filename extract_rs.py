import docx
d = docx.Document(r'D:\Documents\PA 111\source-materials\PA 111-Reading Set 1 (1).docx')
text = []
for p in d.paragraphs:
    text.append(p.text)
with open(r'D:\Documents\PA 111\docs\READINGSET1_EXTRACT.txt','w',encoding='utf-8') as f:
    f.write('\n\n'.join(text))
print(len(d.paragraphs))
